from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT = r"E:\boxeo-ssjj\Manual_Tecnico_Maestro_Combate_Boxeo_Godot_3_Fases_COMPLETO.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text, bold=False, color="000000"):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(str(text))
    run.bold = bold
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def style_table(table, widths=None, header_fill="1F4E79"):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    header_tr_pr = table.rows[0]._tr.get_or_add_trPr()
    if header_tr_pr.find(qn("w:tblHeader")) is None:
        header_tr_pr.append(OxmlElement("w:tblHeader"))
    for row_idx, row in enumerate(table.rows):
        for col_idx, cell in enumerate(row.cells):
            if widths and col_idx < len(widths):
                cell.width = Cm(widths[col_idx])
            tc_pr = cell._tc.get_or_add_tcPr()
            margins = tc_pr.find(qn("w:tcMar"))
            if margins is None:
                margins = OxmlElement("w:tcMar")
                tc_pr.append(margins)
            for edge in ["top", "start", "bottom", "end"]:
                node = margins.find(qn(f"w:{edge}"))
                if node is None:
                    node = OxmlElement(f"w:{edge}")
                    margins.append(node)
                node.set(qn("w:w"), "90")
                node.set(qn("w:type"), "dxa")
            if row_idx == 0:
                set_cell_shading(cell, header_fill)
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
            elif row_idx % 2 == 0:
                set_cell_shading(cell, "F4F7FA")


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)
        run.font.name = "Aptos Display" if level == 1 else "Aptos"
        run._element.rPr.rFonts.set(qn("w:ascii"), run.font.name)
        run._element.rPr.rFonts.set(qn("w:hAnsi"), run.font.name)
    return p


def add_para(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.08
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        r.bold = True
        r.font.name = "Aptos"
        rest = text[len(bold_lead):]
        p.add_run(rest)
    else:
        p.add_run(text)
    for run in p.runs:
        run.font.name = "Aptos"
        run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
        run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
        run.font.size = Pt(10.5)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)
        for run in p.runs:
            run.font.name = "Aptos"
            run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
            run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
            run.font.size = Pt(10)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)
        for run in p.runs:
            run.font.name = "Aptos"
            run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
            run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
            run.font.size = Pt(10)


def add_table(doc, headers, rows, widths=None, header_fill="1F4E79"):
    table = doc.add_table(rows=1, cols=len(headers))
    for idx, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[idx], h, bold=True, color="FFFFFF")
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], value)
    style_table(table, widths, header_fill)
    doc.add_paragraph()
    return table


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    styles["Normal"].font.size = Pt(10.5)
    for name, size in [("Title", 22), ("Heading 1", 15), ("Heading 2", 12), ("Heading 3", 11)]:
        style = styles[name]
        style.font.name = "Aptos Display" if name in ["Title", "Heading 1"] else "Aptos"
        style._element.rPr.rFonts.set(qn("w:ascii"), style.font.name)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), style.font.name)
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)


def build():
    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Manual Tecnico Maestro Combate Boxeo Godot")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Juego de boxeo 3D en Godot 4 organizado en 3 fases de produccion").bold = True
    add_para(doc, "Este manual consolida la arquitectura, los sistemas de combate y la ruta de ejecucion para llevar el juego desde el vertical slice actual hasta un combate tecnico, expresivo y listo para produccion. Esta version esta organizada en tres fases claras: base jugable, combate maestro y produccion final.")
    add_table(doc, ["Campo", "Definicion"], [
        ["Proyecto", "Boxeo 3D en Godot 4 con Quick Fight, seleccion de peleadores, IA, HUD, referee, telemetria, rounds y pruebas automatizadas."],
        ["Objetivo", "Completar un sistema de combate hibrido donde las reglas sigan siendo deterministas y la presentacion exprese peso, direccion, distancia, impacto y recuperacion."],
        ["Regla central", "FightManager, BoxerController, CombatRules, MoveData, Events y los runners existentes se preservan; la nueva presentacion no debe romper el sim ni la jugabilidad actual."],
    ], widths=[4.2, 12.5])

    add_heading(doc, "Indice Operativo", 1)
    add_table(doc, ["Parte", "Contenido"], [
        ["Parte 1", "Vision, estado actual, arquitectura y contratos de datos."],
        ["Parte 2", "Sistema de combate: input, footwork, golpes, defensa, impacto, clinch, IA y presentacion."],
        ["Parte 3", "Ruta de ejecucion en 3 fases, criterios de aceptacion, pruebas, riesgos y prompts de trabajo."],
        ["Apendices", "Tablas maestras de golpes, animaciones, telemetria, runners y checklist final."],
    ], widths=[3.2, 13.5])

    add_heading(doc, "Parte 1 Vision Y Arquitectura", 1)
    add_heading(doc, "1 Vision del juego", 2)
    add_para(doc, "El juego debe sentirse como boxeo profesional: distancia, timing, guardia, stamina, estabilidad, counters, knockdowns y decision por jueces deben importar tanto como acertar golpes. La prioridad no es agregar contenido de MMA ni ligas reales, sino consolidar una experiencia de boxeo 3D legible, reactiva y pulida.")
    add_para(doc, "El estado actual ya contiene una base jugable importante: FightScene instancia ring, peleadores, camara, HUD, audio, referee, FightManager e HitFeedbackSystem; FightManager controla rounds, campana, knockdowns, KO, TKO, decision, estadisticas y eventos globales; BoxerController concentra movimiento, ataques, defensa, clinch, stamina, estabilidad, IA y animaciones.")
    add_heading(doc, "2 Principios rectores", 2)
    add_table(doc, ["Principio", "Implicacion"], [
        ["Reglas antes que espectaculo", "El resultado del combate lo deciden CombatRules, FightManager y BoxerController. La presentacion nunca cambia dano, stamina, hitboxes logicos ni resultado."],
        ["Lectura clara", "Cada golpe debe poder leerse por startup, mano usada, altura, distancia y recovery. El jugador debe entender por que conecto, fallo o fue bloqueado."],
        ["Migracion gradual", "No se reemplazan todos los clips a la vez. Cada golpe puede usar una mezcla de clip y procedural mediante un peso de migracion."],
        ["Un evento, muchos consumidores", "Events publica hechos del combate para HUD, audio, camara, VFX, telemetria y futuros replays sin acoplarlos al interior del peleador."],
        ["Pruebas por fase", "Cada fase termina con runners en PASS, build jugable y lista de riesgos. No se avanza con el juego roto."],
    ], widths=[4.4, 12.3])

    add_heading(doc, "3 Arquitectura actual del combate", 2)
    add_table(doc, ["Sistema", "Rol actual", "Archivo principal"], [
        ["FightScene", "Construye la escena de pelea, selecciona peleadores, configura camara, HUD, audio, referee, feedback y manager.", "scripts/fight/fight_scene.gd"],
        ["FightManager", "Administra round, reloj, campana, knockdown, KO/TKO, decision, scorecards, estadisticas y emision de eventos.", "scripts/managers/fight_manager.gd"],
        ["BoxerController", "Autoridad de cada peleador: movimiento, ataque, defensa, stamina, estabilidad, clinch, IA, animacion y hurtboxes.", "fighters/shared/boxer_controller.gd"],
        ["CombatRules", "Datos y calculo base de ataques, dano, stamina, stun, estabilidad y severidad.", "scripts/combat/combat_rules.gd"],
        ["FootworkModel", "Clasifica distancia, base de movimiento relativo, separacion y reglas de desplazamiento.", "scripts/combat/boxing_footwork_model.gd"],
        ["Events", "Canal global para fight_started, round_started, punch_landed, knockdowns y final de pelea.", "autoload/events.gd"],
    ], widths=[3.2, 9.2, 4.5])
    add_para(doc, "La arquitectura recomendada para las 3 fases es de estrangulamiento: primero estabilizar lo existente, luego extraer contratos de presentacion y finalmente mover la expresion visual avanzada a modulos especializados. BoxerController puede seguir siendo la autoridad inicial, pero cada fase debe reducir dependencias internas y mover responsabilidad a componentes claros.")

    add_heading(doc, "4 Contrato entre simulacion y presentacion", 2)
    add_para(doc, "La presentacion avanzada debe leer snapshots del combate. Un snapshot contiene posicion, velocidad, direccion, estado de ataque, frame de fase, guardia, stamina, estabilidad, stun, clinch y datos del rival. Los eventos de impacto incluyen atacante, defensor, golpe, zona, bloqueo, counter, limpieza, potencia normalizada y punto de contacto.")
    add_table(doc, ["Dato", "Fuente", "Consumidores"], [
        ["FighterView", "BoxerController o PresentationBridge", "Motor procedural, overlay, camara, telemetria visual."],
        ["CombatEvent", "FightManager y Events", "ImpactSolver, HitFeedbackSystem, audio, HUD, VFX."],
        ["MoveData", "data/moves/*.tres y CombatRules", "Timing, coste, rango, objetivo, clips y parametros procedurales."],
        ["Animation Catalog", "docs/ANIMATION_CATALOG.md", "Validacion de clips por peleador y fallback de animacion."],
    ], widths=[3.6, 5.6, 7.5])

    add_heading(doc, "Parte 2 Sistemas Del Juego", 1)
    add_heading(doc, "5 Flujo completo de una pelea", 2)
    add_numbered(doc, [
        "Main menu define modo, peleador, rival, dificultad, rounds y duracion.",
        "FightScene carga ring, peleadores, camara, HUD, audio, referee y manager.",
        "FightManager inicia round, habilita pelea y emite fight_started y round_started.",
        "Cada BoxerController procesa input o IA, actualiza rango, footwork, defensa, ataque, stamina, estabilidad y animacion.",
        "Los golpes activan fist hitboxes, consultan hurtboxes, aplican CombatRules y emiten punch_landed o eventos de bloqueo/counter.",
        "FightStats, Judges, HUD, audio, camara y telemetria consumen el resultado.",
        "Knockdown pausa la pelea, referee cuenta, decide recuperacion, KO o TKO.",
        "Al terminar los rounds, Judges decide y HUD muestra resultado con estadisticas.",
    ])

    add_heading(doc, "6 Input, controles y buffer", 2)
    add_para(doc, "El input debe priorizar respuesta rapida y lectura tactica. El movimiento usa ejes relativos al rival; los ataques deben entrar por acciones discretas o gestos; la defensa debe permitir bloqueo alto, bloqueo al cuerpo, slips, duck, pivots y clinch segun distancia.")
    add_table(doc, ["Accion", "Uso", "Regla tecnica"], [
        ["Movimiento", "move_left, move_right, move_forward, move_backward", "Se interpreta relativo al rival, con deadzone radial y transicion short, medium o long step."],
        ["Jab", "Ataque rapido de control", "Bajo coste, menor dano, activa counters y mide distancia."],
        ["Cross", "Golpe recto de potencia", "Mayor recovery y compromiso de peso; castiga whiffs y guardias abiertas."],
        ["Hooks", "Golpes laterales", "Eficaces en pocket y contra guardia mal orientada."],
        ["Uppercut", "Golpe ascendente", "Eficaz a corta distancia y contra duck o guardia baja."],
        ["Body modifier", "Convierte golpe a cuerpo", "Debe afectar stamina, fatiga y guardia mas que stun de cabeza."],
        ["Defensa", "Bloqueos, slips, duck, pivot", "Consume guardia o stamina; abre counter_window si se temporiza bien."],
        ["Clinch", "Corta distancia y supervivencia", "Solo en pocket o too close; recupera stamina y fuerza separacion posterior."],
    ], widths=[3, 5.2, 8.5])

    add_heading(doc, "7 Footwork y distancia", 2)
    add_para(doc, "El footwork define el combate. Los rangos recomendados son outside, long, mid, pocket y too close. La IA, el HUD de debug y las decisiones de ataque deben depender de esta clasificacion. La separacion corporal evita interpenetracion, pero debe ser suave para no romper la sensacion de peso.")
    add_table(doc, ["Rango", "Distancia de partida", "Conducta esperada"], [
        ["Outside", "Mayor a 2.80 m", "Aproximacion, jab al aire no debe conectar, camara abre encuadre."],
        ["Long", "2.10 m a 2.80 m", "Jab y cross largos; pasos de entrada y salida."],
        ["Mid", "1.35 m a 2.10 m", "Boxeo principal, combinaciones y counters."],
        ["Pocket", "0.78 m a 1.35 m", "Hooks, uppercuts, cuerpo, slips y riesgo de clinch."],
        ["Too close", "Menor a 0.78 m", "Separacion, clinch, bloqueo y micro pasos; golpes largos pierden eficacia."],
    ], widths=[3, 4.2, 9.5])
    add_bullets(doc, [
        "El pie mas cercano a la direccion se mueve primero: delantero para avanzar, trasero para retroceder y lateral del lado para desplazamiento lateral.",
        "El pivot debe poner al rival en desventaja posicional durante una ventana corta, ya implementada como positional_disadvantage_frames.",
        "La fatiga reduce aceleracion, velocidad de recovery, guardia y estabilidad, no solo dano.",
    ])

    add_heading(doc, "8 Golpes y datos maestros", 2)
    add_para(doc, "Cada golpe debe tener timing, rango, objetivo, mano, coste, dano, stun, estabilidad, hitstop y clip de respaldo. La ampliacion procedural agrega trayectoria, rotacion, transferencia de peso, objetivo corporal y parametros de impacto sin eliminar los clips.")
    add_table(doc, ["Golpe", "Rol", "Timing inicial", "Riesgo"], [
        ["Jab", "Medir distancia, interrumpir y preparar combinacion", "startup 4, active 2, recovery 10", "Bajo dano si se abusa; counter si queda corto."],
        ["Cross", "Golpe recto fuerte", "startup 6, active 3, recovery 14", "Recovery mayor y compromiso de peso."],
        ["Left hook", "Castigo lateral en pocket", "startup 7, active 3, recovery 16", "Falla fuera de rango y abre counter."],
        ["Right hook", "Potencia lateral", "startup 8, active 3, recovery 17", "Alta recompensa, alta exposicion."],
        ["Uppercut", "Castigo vertical a corta distancia", "startup 7, active 3, recovery 16", "Pierde contra distancia larga o pivot."],
        ["Body variants", "Drenar stamina y bajar guardia", "Mismo golpe con target body", "Exponen cabeza si se leen mal."],
    ], widths=[3, 5.1, 4.2, 4.4])

    add_heading(doc, "9 Defensa, guardia y counters", 2)
    add_para(doc, "La defensa debe tener coste, timing y consecuencia. Bloquear protege pero consume guard_stamina; slips y ducks no son invulnerabilidad gratis, sino movimientos con entrada y salida. Un counter se reconoce cuando el atacante rival esta en startup o recovery y el golpe propio conecta limpio.")
    add_table(doc, ["Defensa", "Ventaja", "Coste o riesgo"], [
        ["High block", "Reduce golpes a cabeza, permite aguantar presion.", "Consume guard_stamina; vulnerable a cuerpo y guard break."],
        ["Body block", "Protege stamina y torso.", "Abre cabeza si se predice mal."],
        ["Slip left right", "Evita rectos y habilita counter.", "Pierde contra hooks o golpes al cuerpo si se spamea."],
        ["Duck", "Evita hooks altos y prepara uppercut.", "Vulnerable a uppercut y golpes al cuerpo."],
        ["Pivot", "Cambia angulo y genera desventaja posicional.", "Tiene cooldown y requiere rango suficiente."],
        ["Clinch", "Corta combo y recupera parcialmente.", "Debe romperse por referee; no puede ser infinito."],
    ], widths=[3.2, 6.2, 7.3])

    add_heading(doc, "10 Impacto, reaccion y feedback", 2)
    add_para(doc, "El impacto debe sentirse por tres capas coordinadas: resultado logico, reaccion corporal y feedback audiovisual. CombatRules decide dano, stamina, stun, estabilidad y severidad. ImpactSolver visual debe traducir direccion, zona, bloqueo, counter y limpieza en head snap, torso, pushback, hitstop, camara, audio, VFX y vibracion.")
    add_table(doc, ["Evento", "Respuesta visual", "Respuesta de sistema"], [
        ["Jab limpio a cabeza", "Snap corto de cabeza y pequeno hitstop.", "Punch landed, dano bajo, posible interrupcion."],
        ["Cross counter", "Cabeza y torso reaccionan; camara y audio mas fuertes.", "Mayor stun, estabilidad y posible wobble."],
        ["Golpe bloqueado", "Impacto en guardia, brazos absorben.", "Dano reducido, guard_stamina baja, puede romper guardia."],
        ["Golpe al cuerpo", "Torso se cierra, respiracion y fatiga visibles.", "Stamina, body_health y long_term_fatigue bajan."],
        ["Knockdown", "Animacion o ragdoll controlado; referee entra.", "FightManager cambia a KNOCKDOWN y ejecuta conteo."],
    ], widths=[4, 6.7, 6])

    add_heading(doc, "11 IA y tactica", 2)
    add_para(doc, "La IA debe parecer boxeador, no perseguidor. Debe alternar aproximacion, control de rango, defensa, counters, supervivencia, castigo al spam de jab, presion cuando el rival esta wobble y salida de cuerdas.")
    add_table(doc, ["Modo IA", "Condicion", "Conducta"], [
        ["BOX", "Estado normal", "Circula, mide distancia, ataca segun rango y stamina."],
        ["SURVIVE", "Vida, stamina o estabilidad baja", "Retrocede, bloquea, clinch si esta cerca."],
        ["FINISH", "Rival WOBBLED", "Presiona con combinaciones cortas y evita whiffs largos."],
        ["ANTI_JAB_SPAM", "Memoria detecta jab repetido", "Slip, pivot o counter con cross/hook."],
        ["ANTI_KITING", "Rival muy lejos", "Corta ring y evita perseguir en linea recta."],
        ["ANTI_BLOCK", "Rival bloqueando", "Cambia a cuerpo, hook lateral o espera guard break."],
    ], widths=[3.4, 5.4, 7.9])

    add_heading(doc, "12 Animaciones y rig", 2)
    add_para(doc, "Cada peleador seleccionable debe exponer el set minimo de clips de Boxing y UnarmedSupport. El AnimationTree debe iniciar activo en Footwork. Root motion queda desactivado en locomocion compartida: la traslacion pertenece al CharacterBody3D y a BoxerController.")
    add_table(doc, ["Categoria", "Clips obligatorios"], [
        ["Base", "boxing_idle, step_forward, step_backward, step_left, step_right, get_up"],
        ["Golpes", "jab, cross, left_hook, right_hook, uppercut"],
        ["Defensa", "block_left, block_right, block_body"],
        ["Soporte", "block, block_get_hit_1, block_get_hit_2, dodge_backward, dodge_left, dodge_right"],
        ["Reacciones", "get_hit_back, get_hit_front, get_hit_left, get_hit_right, stunned, knockdown, get_up"],
    ], widths=[3.5, 13.2])

    add_heading(doc, "13 Telemetria, HUD y herramientas", 2)
    add_bullets(doc, [
        "HUD de debug de footwork: rango, distancia, intensidad, input, velocidad, target, alineacion, separacion e IA.",
        "HUD de debug de golpes: ataque, fase, mano, rango, target, hitbox activa, resultado, dano, coste, counter, distancia, IA y referee.",
        "Telemetria por pelea: punches thrown, landed, accuracy, knockdowns, scorecards y resultado.",
        "Overlay futuro: targets de mano, COM, poligono de apoyo, vector de impacto, frame activa, hitstop y estado de pies.",
        "Grabacion de sesiones: inputs, semilla, eventos y snapshots para repetir bugs de combate.",
    ])

    add_heading(doc, "Parte 3 Ejecucion En 3 Fases", 1)
    add_heading(doc, "14 Resumen de fases", 2)
    add_table(doc, ["Fase", "Nombre", "Objetivo", "Puerta de salida"], [
        ["1", "Base jugable solida", "Dejar el Quick Fight actual estable, verificable, selectable y con telemetria confiable.", "Runners base en PASS, tres peleadores seleccionables, rounds y resultados correctos."],
        ["2", "Combate maestro", "Mejorar profundidad tactica: datos de golpes, defensa, counters, footwork, IA y presentacion de impacto.", "Cada golpe se lee bien, defensa tiene coste, IA usa rangos y counters, feedback sincronizado."],
        ["3", "Produccion final", "Pulir experiencia completa: carrera ligera, settings, accesibilidad, optimizacion, replays, QA y build final.", "Build jugable estable, checklist completo, sin regresiones criticas."],
    ], widths=[1.5, 4.2, 7.8, 5.2])

    add_heading(doc, "15 Fase 1 Base jugable solida", 2)
    add_para(doc, "Objetivo: asegurar que el juego actual sea una vertical slice confiable. La fase no busca espectacularidad procedural; busca eliminar fragilidad en flujo, seleccion, combate basico, UI, guardado y pruebas.")
    add_table(doc, ["Area", "Entregable", "Criterio"], [
        ["Flujo", "Main menu, fighter select, settings y fight scene conectados.", "Se puede iniciar, reiniciar, salir y cambiar settings sin romper escena."],
        ["Peleadores", "boxer_green, boxer_02 y boxer_03 seleccionables.", "Cada uno carga modelo, skeleton, clips y datos sin errores bloqueantes."],
        ["Combate", "Ataques basicos, bloqueo, stamina, estabilidad, knockdown, KO/TKO y decision.", "Una pelea completa termina por KO, TKO o decision."],
        ["Managers", "FightManager, RoundManager, FightStats, Judges y Telemetry integrados.", "HUD y resultado muestran round, tiempo, ganador, metodo y stats."],
        ["Eventos", "Events emite hechos centrales una vez.", "HUD, audio y feedback no dependen de acceder a internals innecesarios."],
        ["Pruebas", "Runners de lifecycle, combat, live_flow, settings y animation runtime.", "PASS documentado con warnings conocidos separados de fallos."],
    ], widths=[3, 6.5, 7.2])
    add_heading(doc, "Checklist fase 1", 3)
    add_bullets(doc, [
        "El jugador puede completar tres rounds sin bloqueo.",
        "El referee entra en knockdown, cuenta y reanuda o termina pelea.",
        "Los scorecards se guardan y la pantalla final muestra datos coherentes.",
        "La IA pelea en Easy, Medium y Hard sin quedarse inmovil.",
        "Los logs distinguen fallos reales de warnings conocidos.",
    ])

    add_heading(doc, "16 Fase 2 Combate maestro", 2)
    add_para(doc, "Objetivo: convertir la base en boxeo tecnico. Aqui se trabaja la sensacion: timing, lectura, peso del golpe, defensa util, counters, IA tactica y feedback audiovisual. Esta fase debe hacerse por golpes y por sistemas pequenos, no como reescritura total.")
    add_table(doc, ["Bloque", "Trabajo", "Criterio de aceptacion"], [
        ["MoveData ampliado", "Pasar datos de ataques a recursos con timing, rango, target, coste y parametros visuales.", "Cambiar un valor de jab no exige tocar BoxerController."],
        ["Golpes", "Afinar jab, cross, hooks, uppercut y variantes al cuerpo.", "Contacto visual y logico con desfase maximo de 1 frame."],
        ["Defensa", "Costes y ventanas para high block, body block, slip, duck, pivot y clinch.", "Cada defensa tiene counterplay claro."],
        ["Impacto", "ImpactSolver visual para direccion, zona, bloqueo, counter, limpieza y potencia.", "Golpes distintos producen reacciones distintas."],
        ["Footwork", "Micro pasos, pivot, control de distancia y anti foot slide.", "El peleador mantiene stance y no patina visualmente."],
        ["IA", "Modos tacticos y memoria de patrones.", "La IA castiga spam de jab, se protege herida y presiona rival wobble."],
        ["Feedback", "Camara, audio, VFX, HUD y vibracion desde el mismo evento.", "No hay desfase perceptible entre golpe y feedback."],
    ], widths=[3.3, 7.1, 6.3])
    add_heading(doc, "Regla de implementacion fase 2", 3)
    add_para(doc, "Implementar un golpe o sistema por ciclo. Cada ciclo necesita runner rojo primero, cambio minimo, runner verde, revision visual, metrica y nota de riesgo. Si un cambio procedural no mejora al clip, queda detras de toggle o con menor peso de mezcla.")

    add_heading(doc, "17 Fase 3 Produccion final", 2)
    add_para(doc, "Objetivo: convertir el combate en producto completo. Esta fase integra carrera ligera, progresion, accesibilidad, menus finales, optimizacion, replay o modo practica, QA, balance y build.")
    add_table(doc, ["Area", "Entregable", "Criterio"], [
        ["Carrera ligera", "Creacion de personaje, seleccion, historial y progresion basica.", "El jugador puede crear boxeador, pelear y registrar resultado."],
        ["Practica", "Dummy configurable, overlay y repeticion de situaciones.", "Se puede repetir el mismo golpe y distancia para afinado."],
        ["Accesibilidad", "Remapeo, dificultad, ayudas visuales, audio, vibracion y opciones de camara.", "Settings se aplican y persisten."],
        ["Optimizacion", "Perfilado de combate, animacion, fisica, UI y memoria.", "FPS objetivo estable en escena de pelea."],
        ["QA", "Matriz de pruebas por modo, pelea larga, knockdowns, pausa, settings y seleccion.", "Sin regresiones bloqueantes antes de build."],
        ["Build", "Export preset, versionado, changelog y checklist de entrega.", "Build reproducible con resultado documentado."],
    ], widths=[3.1, 7, 6.6])

    add_heading(doc, "18 Runners y puertas", 2)
    add_table(doc, ["Runner", "Comprueba", "Fase"], [
        ["fight_lifecycle_runner.gd", "Inicio, rounds, knockdown, fin de pelea.", "1"],
        ["fight_system_runner.gd", "Integracion de sistemas de pelea.", "1"],
        ["combat_runner.gd", "Ataques, dano, defensa y resultados.", "1 y 2"],
        ["boxing_footwork_controller_runner.gd", "Rangos, pasos, pivots y separacion.", "2"],
        ["technical_combat_runner.gd", "Counters, guardia, clinch, balance y reglas avanzadas.", "2"],
        ["events_runner.gd", "Emision global consistente.", "1 y 2"],
        ["telemetry_runner.gd", "Stats y metricas correctas.", "1, 2 y 3"],
        ["settings_application_runner.gd", "Opciones aplicadas y persistentes.", "3"],
        ["fighter_animation_runtime_runner.gd", "AnimationTree, clips obligatorios y fallback.", "1 y 2"],
    ], widths=[5, 8.3, 3.4])

    add_heading(doc, "19 Riesgos y mitigaciones", 2)
    add_table(doc, ["Riesgo", "Impacto", "Mitigacion"], [
        ["BoxerController demasiado grande", "Cambios de combate causan regresiones cruzadas.", "Extraer por interfaces: MoveData, PresentationBridge, AI planner, ImpactSolver y FootworkModel."],
        ["Procedural robotico", "Pierde personalidad y peso.", "Mantener clips como fallback, usar mezcla gradual y pruebas A/B."],
        ["Eventos duplicados", "HUD, audio y telemetria cuentan dos veces.", "Emitir hechos centrales desde FightManager y validar con events_runner."],
        ["Foot slide", "Rompe credibilidad visual.", "Foot plant, velocidades maximas y runner con umbral de pie plantado."],
        ["IA injusta", "Dificultad se siente tramposa.", "Reaction frames por dificultad, telemetria de counters y ventanas visibles."],
        ["Ragdoll rompe sim", "Resultado no determinista.", "Ragdoll solo en presentacion, nunca escribe al estado logico."],
    ], widths=[4.3, 5.6, 6.8])

    add_heading(doc, "20 Prompts operativos por fase", 2)
    add_para(doc, "Prompt maestro: trabaja en el juego de boxeo Godot sin romper el flujo actual. Conserva FightManager, BoxerController, CombatRules, Events, HUD, telemetria, seleccion de peleadores y runners. Implementa cambios pequenos, con runner rojo primero y runner verde antes de pasar al siguiente bloque. Reporta archivos cambiados, pruebas ejecutadas, pass fail, warnings nuevos y warnings conocidos.")
    add_table(doc, ["Fase", "Prompt breve"], [
        ["1", "Estabiliza Quick Fight: seleccion, pelea completa, rounds, referee, KO/TKO/decision, HUD, audio y telemetria. No agregues sistemas avanzados hasta que los runners base pasen."],
        ["2", "Mejora combate tecnico por ciclos: MoveData, jab, cross, hooks, uppercut, defensa, counters, impacto, footwork e IA. Un golpe o sistema por tarea, con toggle y metrica."],
        ["3", "Prepara produccion: carrera ligera, practica, accesibilidad, settings persistentes, optimizacion, QA, export y checklist final. No aceptes contenido nuevo si rompe estabilidad."],
    ], widths=[1.7, 15])

    add_heading(doc, "Apendice A Tabla Maestra De Golpes", 1)
    add_table(doc, ["Golpe", "Mano", "Target", "Rango", "Coste", "Dano", "Stun", "Estabilidad", "Clip"], [
        ["jab", "left", "head", "long mid", "bajo", "bajo", "bajo", "bajo", "Boxing/jab"],
        ["cross", "right", "head", "mid", "medio", "medio alto", "medio", "medio", "Boxing/cross"],
        ["left_hook", "left", "head", "pocket mid", "medio", "medio", "medio", "medio", "Boxing/left_hook"],
        ["right_hook", "right", "head", "pocket", "alto", "alto", "alto", "alto", "Boxing/right_hook"],
        ["uppercut", "variable", "head", "pocket", "alto", "alto", "alto", "alto", "Boxing/uppercut"],
        ["jab_body", "left", "body", "mid", "bajo medio", "bajo", "bajo", "bajo", "Boxing/jab"],
        ["cross_body", "right", "body", "mid pocket", "medio", "medio", "bajo", "medio", "Boxing/cross"],
        ["hook_body", "left/right", "body", "pocket", "medio alto", "medio", "bajo", "medio", "Boxing/left_hook/right_hook"],
    ], widths=[2.5, 1.8, 1.8, 2.5, 1.8, 1.9, 1.6, 2, 2.8])

    add_heading(doc, "Apendice B Checklist Final De Aceptacion", 1)
    add_bullets(doc, [
        "El juego inicia desde menu y llega a pelea sin errores bloqueantes.",
        "Los tres peleadores base son seleccionables y conservan AnimationTree activo en Footwork.",
        "Una pelea termina correctamente por decision, KO y TKO.",
        "Jab, cross, hooks, uppercut y golpes al cuerpo tienen lectura, coste y counterplay.",
        "Bloqueo, slips, duck, pivot y clinch tienen ventanas y riesgos claros.",
        "El HUD muestra informacion correcta y no duplica eventos.",
        "La IA usa rango, stamina, defensa, counters y modo supervivencia.",
        "La telemetria permite comparar builds y balance.",
        "Los runners de fase pasan y los warnings conocidos estan documentados.",
        "La build final incluye configuracion, remapeo, audio, camara y export reproducible.",
    ])

    add_heading(doc, "Apendice C Glosario", 1)
    add_table(doc, ["Termino", "Definicion"], [
        ["FightSim", "Fuente logica de verdad del combate. En el proyecto actual esta distribuida entre BoxerController, CombatRules y managers."],
        ["Presentacion", "Capa visual y audiovisual que expresa hechos ya decididos por la logica."],
        ["Hitstop", "Congelacion breve al contacto para dar peso al impacto."],
        ["Wobble", "Estado de inestabilidad alta donde el peleador queda vulnerable."],
        ["Pocket", "Rango corto donde hooks, uppercuts, cuerpo y clinch dominan."],
        ["Counter", "Golpe que conecta durante startup o recovery del rival, con recompensa aumentada."],
        ["Guard break", "Ruptura temporal de defensa por agotar guard_stamina o recibir impacto fuerte."],
    ], widths=[3.6, 13.1])

    section = doc.sections[0]
    footer = section.footer
    footer_p = footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_p.text = "Manual tecnico maestro de combate boxeo Godot en 3 fases"

    doc.save(OUTPUT)


if __name__ == "__main__":
    build()
